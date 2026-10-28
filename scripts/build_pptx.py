"""Build the proposal slide deck (docs/Proposal_AI_Air_Doctor.pptx)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

BASE = os.path.join(os.path.dirname(__file__), "..")
IMG = os.path.join(BASE, "docs", "img")
OUT = os.path.join(BASE, "docs", "Proposal_AI_Air_Doctor.pptx")
FONT = "Microsoft JhengHei"

BG = RGBColor(0x0B, 0x12, 0x20)
CARD = RGBColor(0x15, 0x20, 0x33)
CARD2 = RGBColor(0x1E, 0x2B, 0x42)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
GREEN = RGBColor(0x34, 0xD3, 0x99)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
RED = RGBColor(0xF8, 0x71, 0x71)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 14
_page = [0]


def _font(run, size, color=WHITE, bold=False, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    name = "Consolas" if mono else FONT
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT if not mono else "Consolas")


def text(slide, x, y, w, h, lines, size=16, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, mono=False, spacing=1.15):
    """lines: str or list of str / (str, dict) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        opts = {}
        if isinstance(ln, tuple):
            ln, opts = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.line_spacing = spacing
        p.space_after = Pt(opts.get("after", 4))
        _font(p.add_run(), 1)  # placeholder removed below
        p.runs[0].text = ln
        _font(p.runs[0], opts.get("size", size), opts.get("color", color),
              opts.get("bold", bold), mono)
    return tb


def rect(slide, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def new_slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    _page[0] += 1
    if title:
        rect(s, 0.6, 0.55, 0.1, 0.6, CYAN, shape=MSO_SHAPE.RECTANGLE)
        text(s, 0.85, 0.42, 11, 0.7, title, 30, WHITE, True)
        if subtitle:
            text(s, 0.87, 1.08, 11, 0.4, subtitle, 14, MUTED)
        text(s, 11.6, 6.95, 1.4, 0.35, f"{_page[0]} / {TOTAL}", 11, MUTED, align=PP_ALIGN.RIGHT)
        text(s, 0.6, 6.95, 6, 0.35, "AI Air Doctor · Final Project Proposal", 11, MUTED)
    return s


def card(slide, x, y, w, h, title, body, accent=CYAN, title_size=17, body_size=13):
    rect(slide, x, y, w, h, CARD)
    rect(slide, x, y, w, 0.07, accent, shape=MSO_SHAPE.RECTANGLE)
    text(slide, x + 0.2, y + 0.2, w - 0.4, 0.5, title, title_size, accent, True)
    text(slide, x + 0.2, y + 0.72, w - 0.4, h - 0.85, body, body_size, WHITE)


def table(slide, x, y, w, rows, col_w, size=12, row_h=0.5, header_fill=RGBColor(0x0E, 0x74, 0x90)):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                 Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j, val in enumerate(r):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (CARD if i % 2 else CARD2)
            c.margin_left = c.margin_right = Inches(0.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = val
            color = WHITE
            if val.startswith("✓"):
                color = GREEN
            elif val.startswith("✗"):
                color = RED
            elif val.startswith("△"):
                color = AMBER
            _font(run, size, color, i == 0)
    return tbl


# ---------------- 1. Title ----------------
s = new_slide(None)
rect(s, 0, 0, 13.333, 7.5, BG, shape=MSO_SHAPE.RECTANGLE)
for i, (c, x, y, d) in enumerate([(CYAN, 9.6, -1.2, 4.5), (VIOLET, 10.9, 3.6, 3.2), (GREEN, 8.4, 4.9, 1.6)]):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    o.fill.solid()
    o.fill.fore_color.rgb = c
    o.fill.transparency = 0.0
    o.line.fill.background()
    # semi-transparent via XML alpha
    sf = o.fill._xPr.find(qn("a:solidFill"))
    clr = sf[0]
    a = clr.makeelement(qn("a:alpha"), {"val": "22000"})
    clr.append(a)
text(s, 0.9, 1.5, 9, 0.5, "FINAL PROJECT PROPOSAL", 16, CYAN, True)
text(s, 0.9, 2.0, 10, 1.2, "AI Air Doctor", 60, WHITE, True)
text(s, 0.9, 3.2, 10, 1.0, "結合 IoT 空品數據與環境醫學知識的\n個人化空氣品質健康代理人", 26, WHITE)
rect(s, 0.9, 4.65, 1.2, 0.06, CYAN, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.9, 4.9, 10, 0.5, "AI Agent  +  Domain Knowledge  +  Data Analytics  +  Real-World Application", 15, MUTED)
text(s, 0.9, 5.8, 10, 0.8, ["組員:________、________、________、________", "課程:IoT ________    指導老師:________"], 14, MUTED)

# ---------------- 2. Motivation ----------------
s = new_slide("研究動機 Motivation", "空氣污染是全球第二大早逝風險因子")
stats = [("790 萬", "2023 年全球\n空污致死人數", CYAN), ("1 / 8", "全球每 8 位死者\n就有 1 位與空污有關", VIOLET),
         ("86%", "空污死亡來自心臟病、\n中風、COPD 等慢性病", AMBER), ("No.2", "僅次於高血壓的\n第二大早逝風險因子", RED)]
for i, (num, desc, c) in enumerate(stats):
    x = 0.7 + i * 3.05
    rect(s, x, 1.75, 2.8, 2.6, CARD)
    text(s, x, 1.95, 2.8, 1.0, num, 44, c, True, PP_ALIGN.CENTER)
    text(s, x + 0.15, 3.05, 2.5, 1.2, desc, 14, WHITE, align=PP_ALIGN.CENTER)
text(s, 0.7, 4.45, 12, 0.4, "資料來源:Health Effects Institute, State of Global Air 2025", 11, MUTED)
rect(s, 0.7, 5.0, 11.9, 1.7, CARD2)
text(s, 1.0, 5.15, 11.4, 1.5, [
    ("台灣情境", {"bold": True, "color": CYAN, "size": 17}),
    "西半部受工業、交通排放與秋冬東北季風帶來的境外污染影響,中南部常出現 PM2.5 紅色警戒;"
    "氣喘患者、長者、孕婦與兒童等敏感族群首當其衝。【待補:2025–2026 台灣空污新聞/統計】"], 15)

# ---------------- 3. Problem ----------------
s = new_slide("問題定義 Problem", "民眾真正想問的,不是「AQI 多少」,而是「我該怎麼做」")
rect(s, 0.7, 1.75, 5.6, 1.6, CARD2, line=VIOLET)
text(s, 0.95, 1.9, 5.2, 0.4, "💬 使用者提問", 15, VIOLET, True)
text(s, 0.95, 2.35, 5.2, 1.0, "「我有氣喘,明天早上 7 點\n想去大安森林公園跑步,可以嗎?」", 19, WHITE, True)
text(s, 0.7, 3.65, 5.6, 0.5, "回答這個問題,需要同時具備:", 16, MUTED)
needs = [("① 即時 / 預測空品數據", CYAN), ("② 氣象條件(風、濕度、雨)", GREEN),
         ("③ 個人健康狀況", AMBER), ("④ 環境醫學知識", RED)]
for i, (n, c) in enumerate(needs):
    rect(s, 0.7, 4.15 + i * 0.62, 5.6, 0.52, CARD)
    rect(s, 0.7, 4.15 + i * 0.62, 0.08, 0.52, c, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.95, 4.2 + i * 0.62, 5.2, 0.45, n, 15, WHITE, anchor=MSO_ANCHOR.MIDDLE)
table(s, 6.8, 1.75, 5.85, [
    ["現有方案", "能做到", "做不到"],
    ["空品 App\n(環境即時通)", "顯示即時數值", "✗ 個人化建議"],
    ["PM2.5 預測研究", "預測未來濃度", "✗ 民眾難以理解"],
    ["一般 LLM\n(ChatGPT)", "自然語言對話", "✗ 無即時數據、幻覺"],
    ["近期空品 Agent 研究", "LLM + 資料分析", "△ 偏政策、少個人化"],
], [1.95, 1.85, 2.05], 13, 0.85)

# ---------------- 4. Related Work ----------------
s = new_slide("文獻回顧 Related Work", "½ AI Technology + ½ Domain Know-how")
table(s, 0.6, 1.65, 12.1, [
    ["Ref", "Year", "Problem", "AI Method", "Findings", "Limitation"],
    ["[9]", "2025", "野火期間空品分析與政策建議", "Instructor-Worker 多代理人 LLM", "自動取得、分析資料並產生政策建議", "偏政策層面,無個人化"],
    ["[10]", "2025", "感測數據民眾難以理解", "LLM + Model Context Protocol", "把原始數據轉成可行動回饋", "未包含預測模型"],
    ["[15]", "2025", "即時 AQI 監測與健康建議", "多代理人(收集/預測/健康/溝通)", "提供主動式建議", "【待查證】"],
    ["[11]", "2025", "台灣 PM2.5 預測", "LSTM / Transformer", "【待補】", "只輸出數字,無決策建議"],
    ["[12]", "2025", "極端高污染事件預測", "Transformer + 重採樣", "改善高濃度預測", "【待補】"],
], [0.7, 0.8, 2.8, 3.0, 2.8, 2.0], 12, 0.62)
card(s, 0.6, 5.55, 5.9, 1.25, "AI Technology", "LSTM / Transformer 時序預測 · XGBoost · SHAP 可解釋 · ReAct Agent · RAG", CYAN, 15, 12)
card(s, 6.8, 5.55, 5.9, 1.25, "Domain Knowledge", "WHO 2021 指引 · 台灣 AQI 六級 · PM2.5 健康影響 · 季風與逆溫", GREEN, 15, 12)

# ---------------- 5. Research Gap ----------------
s = new_slide("研究缺口 Research Gap", "Existing Research → Limitations → Research Gap → Our Project")
steps = [("Existing Research", "PM2.5 預測模型\nLLM 空品代理人", CYAN), ("Limitations", "預測只輸出數字\nAgent 偏政策、少個人化", AMBER),
         ("Research Gap", "缺少整合 IoT + 預測 +\n醫學知識 + 個人檔案的 Agent", RED), ("Our Project", "AI Air Doctor", GREEN)]
for i, (t, b, c) in enumerate(steps):
    x = 0.6 + i * 3.1
    sh = s.shapes.add_shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, Inches(x), Inches(1.75), Inches(3.0), Inches(0.75))
    sh.fill.solid(); sh.fill.fore_color.rgb = c; sh.line.fill.background()
    tf = sh.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run(); r.text = t; _font(r, 15, BG, True)
    text(s, x + 0.1, 2.65, 2.8, 1.0, b, 14, WHITE, align=PP_ALIGN.CENTER)
table(s, 0.6, 3.85, 12.1, [
    ["能力", "空品 App", "PM2.5 預測研究", "一般 LLM", "近期空品 Agent", "AI Air Doctor"],
    ["即時 IoT 數據", "✓", "✓", "✗", "✓", "✓"],
    ["24h 污染預測", "✗", "✓", "✗", "△", "✓"],
    ["環境醫學知識 (RAG)", "✗", "✗", "△ 可能幻覺", "△", "✓"],
    ["個人健康檔案", "✗", "✗", "✗", "△", "✓"],
    ["有依據的自然語言建議", "✗", "✗", "✗", "✓", "✓"],
], [2.6, 1.8, 2.0, 1.9, 1.9, 1.9], 13, 0.48)

# ---------------- 6. Solution overview ----------------
s = new_slide("提出方案 Proposed Solution", "AI Air Doctor — 具備環境醫學專業能力的 Domain AI Agent")
pillars = [("🤖  AI Agent", "LLM (ReAct) 理解需求、規劃步驟、自主呼叫 5 個工具並整合推理", CYAN),
           ("📚  Domain Knowledge", "WHO 2021 指引、台灣 AQI 健康建議、污染物健康影響、敏感族群定義", GREEN),
           ("📈  Data Analytics", "XGBoost / LSTM 預測 24h PM2.5,SHAP 解釋風速、濕度等影響因子", AMBER),
           ("🌏  Real-World Application", "敏感族群健康防護、戶外活動規劃、高污染主動提醒", VIOLET)]
for i, (t, b, c) in enumerate(pillars):
    x = 0.6 + (i % 2) * 6.15
    y = 1.75 + (i // 2) * 2.1
    card(s, x, y, 5.95, 1.9, t, b, c, 19, 15)
rect(s, 0.6, 6.0, 12.1, 0.75, CARD2, line=CYAN)
text(s, 0.8, 6.05, 11.7, 0.65, "核心價值:每一則建議都附上「數據來源 + 預測信心 + 知識出處」,降低 LLM 幻覺",
     16, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# ---------------- 7. Agent design ----------------
s = new_slide("Agent 設計 Agent Design", "Role · Tools · Decision Logic")
rect(s, 0.6, 1.7, 4.6, 5.05, CARD)
text(s, 0.85, 1.85, 4.2, 0.45, "🩺 Agent 角色 (System Prompt)", 16, CYAN, True)
text(s, 0.85, 2.35, 4.2, 2.3, [
    "• 結合大氣科學與環境醫學的空品健康顧問",
    "• 必須先用工具取得數據,不可猜測數值",
    "• 健康建議必須引用知識庫來源",
    "• 對敏感族群採較嚴格標準",
    "• 結尾附上免責聲明"], 13)
text(s, 0.85, 4.6, 4.2, 0.45, "🧠 決策邏輯", 16, GREEN, True)
text(s, 0.85, 5.05, 4.2, 1.7, [
    "1. 解析意圖(地點/時間/活動)",
    "2. 判斷族群 → 選擇風險門檻",
    "3. 即時數據 or 預測模型",
    "4. 風險 = PM2.5 × 活動係數",
    "5. RAG 檢索 → 生成 → 檢查引用"], 13)
table(s, 5.5, 1.7, 7.2, [
    ["Tool", "輸入", "輸出", "來源"],
    ["get_realtime_aq", "location", "PM2.5、O₃、AQI", "環境部 / Civil IoT"],
    ["get_weather", "location", "溫溼度、風速風向、雨量", "中央氣象署"],
    ["forecast_pm25", "station, time", "預測值、信心、主因", "XGBoost/LSTM+SHAP"],
    ["search_health_kb", "risk, group", "Top-k 建議 + 出處", "ChromaDB"],
    ["get_user_profile", "user_id", "年齡、疾病、懷孕", "SQLite"],
], [1.85, 1.35, 2.2, 1.8], 12, 0.62)
table(s, 5.5, 5.6, 7.2, [
    ["族群 (PM2.5 μg/m³)", "🟢 低", "🟡 中", "🟠 高", "🔴 極高"],
    ["一般民眾", "≤30", "31–50", "51–70", ">70"],
    ["敏感族群", "≤15", "16–30", "31–50", ">50"],
], [2.0, 1.3, 1.3, 1.3, 1.3], 12, 0.38)

# ---------------- 8. Architecture ----------------
s = new_slide("系統架構 Architecture", "User → LLM Agent → Tools → Data Sources → Personalized Advice")
rect(s, 0.6, 1.6, 12.1, 5.2, WHITE)
s.shapes.add_picture(os.path.join(IMG, "architecture.png"), Inches(1.15), Inches(1.7), height=Inches(5.0))

# ---------------- 9. Flowchart ----------------
s = new_slide("系統流程 Flowchart", "從提問到有依據的建議")
rect(s, 0.6, 1.55, 4.6, 5.25, WHITE)
s.shapes.add_picture(os.path.join(IMG, "flowchart.png"), Inches(1.0), Inches(1.6), height=Inches(5.15))
rect(s, 5.5, 1.55, 7.2, 5.25, CARD)
text(s, 5.8, 1.7, 6.7, 0.5, "📍 範例:氣喘患者明早晨跑", 18, CYAN, True)
text(s, 5.8, 2.3, 6.7, 4.4, [
    ("1  意圖:大安區 · 明早 07:00 · 跑步", {}),
    ("2  檔案:氣喘 → 敏感族群門檻", {}),
    ("3  未來時段 → 執行預測模型", {}),
    ("     預測 PM2.5 = 38 μg/m³(主因:低風速、高濕度)", {"color": MUTED, "size": 13}),
    ("4  風險 = 38 × 1.5(跑步)= 57 → 🔴 極高", {}),
    ("5  RAG 檢索:WHO 指引、AQI 敏感族群建議", {}),
    ("6  生成回答並檢查引用", {}),
    ("", {}),
    ("💡 「建議改到室內運動,或延到下午風速增強後再出門;", {"color": GREEN, "bold": True}),
    ("     若外出請配戴 N95,並隨身攜帶吸入劑。」", {"color": GREEN, "bold": True}),
    ("     (數值為示意)", {"color": MUTED, "size": 12}),
], 15)

# ---------------- 10. Pseudocode ----------------
s = new_slide("虛擬碼 Pseudocode", "Algorithm 1 — AI Air Doctor Agent")
rect(s, 0.6, 1.55, 12.1, 5.25, RGBColor(0x02, 0x06, 0x17), line=CARD2)
code = """Input : user_query q, user_id u      Output: advice A, risk_level, sources
 1: profile ← GetUserProfile(u)                 // 疾病、年齡、懷孕
 2: intent  ← LLM.Parse(q)                      // location, time, activity
 3: station ← NearestStation(intent.location)
 4: if intent.time ≤ now + 1h then
 5:     pm ← GetRealtimeAQ(station).pm25 ;  wx ← GetWeather(station)
 6: else
 7:     X ← BuildFeatures(station, last 24h AQ + weather)
 8:     pm, conf ← ForecastModel.Predict(X, intent.time)
 9:     factors ← SHAP.Explain(X)               // 主要影響因子
10: end if
11: θ    ← Thresholds(profile.group)            // 一般 / 敏感族群
12: risk ← RiskLevel(pm × ActivityFactor(intent.activity), θ)
13: docs ← RAG.Retrieve(risk, profile, intent.activity, top_k = 3)
14: A    ← LLM.Generate(q, pm, risk, factors, docs, profile)
15: if not HasCitations(A, docs) then A ← LLM.Regenerate("cite sources")
16: return A, risk, sources(docs, station)"""
text(s, 0.9, 1.65, 11.6, 5.1, code.split("\n"), 13, RGBColor(0xA5, 0xF3, 0xFC), mono=True, spacing=1.0)

# ---------------- 11. Data & Model ----------------
s = new_slide("資料與預測模型 Data & Model", "全部使用台灣官方開放資料 · 民生公共物聯網")
table(s, 0.6, 1.65, 7.0, [
    ["資料", "內容", "頻率"],
    ["環境部空品測站", "PM2.5、PM10、O₃、NO₂、AQI", "每小時"],
    ["民生公共物聯網", "微型感測器 PM2.5、溫溼度", "~5 分鐘"],
    ["中央氣象署", "溫度、濕度、風速風向、雨量", "每小時"],
    ["自建感測器(選做)", "ESP32 + PMS5003 + DHT22", "每分鐘"],
], [2.2, 3.3, 1.5], 13, 0.62)
card(s, 7.9, 1.65, 4.8, 3.1, "Algorithm 2 · 預測模型訓練", [
    "• 資料:2022–2025 逐時資料",
    "• 補值:KNN Imputation",
    "• 特徵:lag 1–24h、移動平均、時間、風向量、溫溼度、雨量",
    "• 切分:2022–24 訓練 / 2025 測試",
    "• 模型:Persistence / XGBoost / LSTM"], AMBER, 16, 13)
card(s, 0.6, 5.0, 12.1, 1.75, "技術選型", "LLM: Gemini / GPT / Llama 3   ·   Agent: LangGraph   ·   ML: XGBoost、PyTorch、SHAP   ·   "
     "Vector DB: ChromaDB   ·   Backend: FastAPI   ·   Frontend: Streamlit / LINE Bot", CYAN, 16, 14)

# ---------------- 12. Experiments ----------------
s = new_slide("實驗設計與預期成果 Experiments & Expected Findings")
table(s, 0.6, 1.5, 12.1, [
    ["實驗", "方法", "評估指標"],
    ["E1 預測準確度", "Persistence vs. XGBoost vs. LSTM", "MAE、RMSE、R²(1h / 6h / 24h)"],
    ["E2 Agent 回答品質", "AI Air Doctor vs. 一般 LLM,同一組 30 題", "正確性、依據完整度、幻覺率"],
    ["E3 Case Study", "氣喘患者 · 晨跑者 · 孕婦", "建議合理性、個人化程度"],
    ["E4 使用者測試(選做)", "10 位使用者試用", "SUS 易用性、滿意度"],
], [3.0, 5.0, 4.1], 14, 0.6)
text(s, 0.6, 4.75, 6, 0.5, "🎯 Expected Findings(預期成果,非實驗結果)", 17, GREEN, True)
for i, t in enumerate(["24h PM2.5 預測誤差明顯低於 Persistence Baseline",
                       "Agent 建議正確性與引用依據比例高於一般 LLM,幻覺率下降",
                       "能為敏感族群提供可行動的個人化防護建議"]):
    rect(s, 0.6 + i * 4.07, 5.3, 3.9, 1.45, CARD)
    text(s, 0.8 + i * 4.07, 5.4, 3.5, 1.3, t, 15, WHITE, anchor=MSO_ANCHOR.MIDDLE)

# ---------------- 13. Timeline ----------------
s = new_slide("時程規劃 Timeline")
tl = [("W1–2", "Proposal\n文獻 · API 串接", CYAN), ("W3–4", "資料前處理\n預測模型訓練", GREEN),
      ("W5–6", "知識庫建立\nAgent 工具開發", AMBER), ("W7", "前端整合\nStreamlit / LINE", VIOLET),
      ("W8", "實驗評估\n期末報告", RED)]
rect(s, 0.9, 3.55, 11.5, 0.08, CARD2, shape=MSO_SHAPE.RECTANGLE)
for i, (w, t, c) in enumerate(tl):
    x = 0.9 + i * 2.6
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.85), Inches(3.3), Inches(0.55), Inches(0.55))
    o.fill.solid(); o.fill.fore_color.rgb = c; o.line.fill.background()
    text(s, x, 2.5, 2.3, 0.6, w, 22, c, True, PP_ALIGN.CENTER)
    rect(s, x, 4.2, 2.3, 1.5, CARD)
    text(s, x + 0.1, 4.3, 2.1, 1.3, t, 15, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 6.1, 11.5, 0.5, "GitHub Repository + DESIGN.md 從 Proposal 階段開始持續維護", 14, MUTED, align=PP_ALIGN.CENTER)

# ---------------- 14. References ----------------
s = new_slide("參考文獻 References")
refs = [
    "[1] Health Effects Institute (2025). State of Global Air 2025.",
    "[2] WHO (2021). WHO Global Air Quality Guidelines.",
    "[3] 環境部 環境資料開放平台 data.moenv.gov.tw    [4][5] 民生公共物聯網 ci.taiwan.gov.tw",
    "[6] 中央氣象署 氣象資料開放平臺 opendata.cwa.gov.tw",
    "[7] Yao, S. et al. (2023). ReAct: Synergizing Reasoning and Acting in Language Models. ICLR.",
    "[8] Lewis, P. et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS.",
    "[9] Instructor-Worker LLM System for Policy Recommendation: LA Wildfires 2025. arXiv.【待查證】",
    "[10] LLM-enhanced Air Quality Monitoring Interface via Model Context Protocol (2025). arXiv.【待查證】",
    "[11][12] 【待補:2025 台灣 PM2.5 預測論文】",
    "[13] Chen & Guestrin (2016). XGBoost. KDD.    [14] Lundberg & Lee (2017). SHAP. NeurIPS.",
    "[15] Multi-Agent Framework for Real-Time AQI Monitoring and Personalized Health Advice (2025).【待查證】",
]
text(s, 0.7, 1.5, 12, 4.6, refs, 13, WHITE, spacing=1.2)
text(s, 0.7, 6.15, 12, 0.6, "Thank you!  Q & A", 26, CYAN, True, PP_ALIGN.CENTER)

prs.save(OUT)
print("saved", os.path.abspath(OUT))
